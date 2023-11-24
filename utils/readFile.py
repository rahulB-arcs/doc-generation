import csv
import docx2txt
from docx import Document

ALLOWED_EXTENSIONS = {'csv', 'docx'}

def read_doc(doc_path):
    """
    Extract text from Doc files
    """
    try:
        text = docx2txt.process(doc_path)
        return text
    except Exception as e:
        return {'error' : str(e)}


def read_csv(csv_path):
    """
    Extract Data from csv file
    """
    try:
        csv_data = []
        with open(csv_path, 'r', encoding='latin-1') as csv_file:
            csv_reader = csv.reader(csv_file)
            for row in csv_reader:
                csv_data.append(row)
        return csv_data
        
    except Exception as e:
        return {'error':str(e)}
    

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_placeholders(file_path):
    placeholders = []
    doc = Document(file_path)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            text = run.text
            if "{" in text and "}" in text:
                start = text.find("{")
                end = text.find("}")
                placeholder = text[start + 1 : end]
                placeholders.append(placeholder)
    return placeholders

def extract_columns(file_path):
    columns = []
    with open(file_path, 'r') as csv_file:
        csv_reader = csv.reader(csv_file)
        headers = next(csv_reader)
        columns = headers

    return columns

def match_columns(filename, columns, placeholders):
    matched_columns = []
    unmatched_columns = []

    for column in columns:
        if column in placeholders:
            matched_columns.append(column)
        else:
            unmatched_columns.append(column)

    return {
        filename: {
            'matched_columns': matched_columns,
            'unmatched_columns': unmatched_columns
        }
    }

def match_placeholders(filename, columns, placeholders):
    matched_placeholders = []
    unmatched_placeholders = []
    # doc_data = {
    #     "matched_placeholders": [], 
    #     "unmatched_placeholders": []
    # }

    for placeholder in placeholders:
        if placeholder in columns:
            matched_placeholders.append(placeholder)
        else:
            unmatched_placeholders.append(placeholder)
    
    # matched_placeholders = [ph for ph in placeholders if ph in columns]
    # unmatched_placeholders = [ph for ph in placeholders if ph not in columns]

    # doc_data["matched_placeholders"] = matched_placeholders
    # doc_data["unmatched_placeholders"] = unmatched_placeholders

    return {
        filename: {
            'matched_placeholders': matched_placeholders,
            'unmatched_placeholders': unmatched_placeholders
        }
    }

    # return [
    #     f'{filename}: [{doc_data}]'
    # ]