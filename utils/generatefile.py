from utils.readfiles import *
from docx import Document

def genratefile(file_path_doc,file_path_csv,new_docx_file_name):

    """
    generate .docx file  with CSV data
    """

    outputdocpaths=[]
    
    text=read_doc(file_path_doc)
    csvdata=read_csv(file_path_csv)
    headers=csvdata[0]
    for row in range(0,(len(csvdata)-1)):
        row=row+1
    
        values = csvdata[row]
        doc = Document()
        # Create a dictionary with headers as keys and corresponding values
        placeholders = {header: value for header, value in zip(headers, values)}
        filled_template = text.format(**placeholders)
        doc.add_paragraph(filled_template)
        new_docx_file_path='static/genrated_docs/'+new_docx_file_name+'_'+str(row)+'_'+'.docx'
        doc.save(new_docx_file_path)
        outputdocpaths.append(new_docx_file_path)
    return outputdocpaths
    


    

    

# file_path_doc='Dummy.docx'
# file_path_csv='dummy-csv.csv'
# newdoc_heading='Test-doc'
# new_docx_file_path="new_document.docx"
# genratefile(file_path_doc,file_path_csv,new_docx_file_path,newdoc_heading)